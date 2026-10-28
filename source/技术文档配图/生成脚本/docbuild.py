"""Helpers that write into the competition template with its look (宋体 + Times New Roman, 1.5 line spacing)."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RED = 'FF0000'
ALIGN = dict(left=WD_ALIGN_PARAGRAPH.LEFT, center=WD_ALIGN_PARAGRAPH.CENTER,
             justify=WD_ALIGN_PARAGRAPH.JUSTIFY, right=WD_ALIGN_PARAGRAPH.RIGHT)


class Builder:
    TEXT_W = 14.6  # cm, A4 with the template's 3.17 cm side margins
    def __init__(self, template):
        self.doc = Document(template)
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn('w:sectPr'):
                body.remove(child)
        self.part = 0
        self.nfig = 0
        self.ntab = 0
        self.counts = {}
        self._count_key = None

    # ------------------------------------------------------------------ low level
    @staticmethod
    def font(run, size=12, bold=False, color=None, east='宋体', west='Times New Roman', italic=False):
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rFonts.set(qn(k), west)
        rFonts.set(qn('w:eastAsia'), east)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def runs(self, p, text, size=12, bold=False, color=None):
        """**bold** and 【placeholder】 (red) inline markup."""
        for seg in re.split(r'(\*\*.+?\*\*|【[^】]*】)', text):
            if not seg:
                continue
            b, c, t = bold, color, seg
            if seg.startswith('**') and seg.endswith('**'):
                b, t = True, seg[2:-2]
            elif seg.startswith('【'):
                c = RED
            r = p.add_run(t)
            self.font(r, size=size, bold=b, color=c)
        if self._count_key:
            plain = re.sub(r'\*\*|\s', '', text)
            self.counts[self._count_key] = self.counts.get(self._count_key, 0) + len(plain)

    @staticmethod
    def _nogrid(p):
        pPr = p._p.get_or_add_pPr()
        el = OxmlElement('w:snapToGrid')
        el.set(qn('w:val'), '0')
        pPr.insert_element_before(el, 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
                                  'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
                                  'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr',
                                  'w:sectPr', 'w:pPrChange')

    @staticmethod
    def _outline(p, lvl):
        pPr = p._p.get_or_add_pPr()
        el = OxmlElement('w:outlineLvl')
        el.set(qn('w:val'), str(lvl))
        pPr.insert_element_before(el, 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')

    def para(self, text='', size=12, bold=False, align='justify', indent=True, before=0, after=0, line=1.5,
             outline=None, keep_next=False, color=None, left_chars=0):
        p = self.doc.add_paragraph()
        self._nogrid(p)
        pf = p.paragraph_format
        pf.line_spacing = line
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = keep_next
        p.alignment = ALIGN[align]
        if indent or left_chars:
            ind = p._p.get_or_add_pPr().get_or_add_ind()
            if indent:
                ind.set(qn('w:firstLineChars'), '200')
                ind.set(qn('w:firstLine'), str(int(size * 20 * 2)))
            if left_chars:
                ind.set(qn('w:leftChars'), str(left_chars * 100))
                ind.set(qn('w:left'), str(int(size * 20 * left_chars)))
        if outline is not None:
            self._outline(p, outline)
        if text:
            self.runs(p, text, size=size, bold=bold, color=color)
        return p

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ structure
    def title(self, text):
        try:
            st = self.doc.styles['Heading 2']
        except KeyError:
            st = None
        p = self.doc.add_paragraph(style=st)
        p.alignment = ALIGN['center']
        p.paragraph_format.line_spacing = 1.5
        for i, part in enumerate(text.split('\n')):
            r = p.add_run(part)
            self.font(r, size=18, bold=True, east='黑体')
            if i < len(text.split('\n')) - 1:
                r.add_break()
        return p

    def part_heading(self, text, new_page=True):
        if new_page:
            self.page_break()
        self.part += 1
        self.nfig = self.ntab = 0
        self._count_key = None
        return self.para(text, size=14, bold=True, indent=False, before=6, after=6, outline=0, keep_next=True)

    def h2(self, text, limit_key=None):
        self._count_key = limit_key
        p = self.para(text, size=12, bold=True, indent=False, before=6, after=3, outline=1, keep_next=True)
        if limit_key:
            self.counts[limit_key] = 0
        return p

    def h3(self, text):
        self._count_key = None
        return self.para(text, size=12, bold=True, indent=False, before=3, after=2, outline=2, keep_next=True)

    def h4(self, text):
        return self.para(text, size=12, bold=True, indent=True, keep_next=True)

    def p(self, text, **kw):
        return self.para(text, **kw)

    def items(self, lines, **kw):
        for t in lines:
            self.para(t, **kw)

    def stop_count(self):
        self._count_key = None

    # ------------------------------------------------------------------ figures
    def fig_caption(self, caption):
        self.nfig += 1
        save = self._count_key
        self._count_key = None
        p = self.para(f'图{self.part}-{self.nfig}  {caption}', size=10.5, align='center', indent=False, after=6, line=1.25)
        self._count_key = save
        return f'图{self.part}-{self.nfig}'

    def next_fig(self, k=1):
        return f'图{self.part}-{self.nfig + k}'

    def next_tab(self, k=1):
        return f'表{self.part}-{self.ntab + k}'

    def figure(self, path, width_cm, caption):
        p = self.doc.add_paragraph()
        p.alignment = ALIGN['center']
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        self._nogrid(p)
        p.add_run().add_picture(path, width=Cm(min(width_cm, self.TEXT_W)))
        return self.fig_caption(caption)

    def placeholder(self, desc, caption, height_cm=5.0, width_cm=13.0):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._borders(t, val='dashed', color='999999')
        cell = t.cell(0, 0)
        cell.width = Cm(width_cm)
        tr = t.rows[0]._tr
        trPr = tr.get_or_add_trPr()
        h = OxmlElement('w:trHeight')
        h.set(qn('w:val'), str(int(height_cm * 567)))
        h.set(qn('w:hRule'), 'exact')
        trPr.append(h)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cp = cell.paragraphs[0]
        cp.alignment = ALIGN['center']
        r = cp.add_run(f'【此处插入：{desc}】')
        self.font(r, size=10.5, color=RED)
        keep = self.doc.add_paragraph()
        keep.paragraph_format.space_after = Pt(0)
        keep.paragraph_format.line_spacing = 0.5
        return self.fig_caption(caption)

    # ------------------------------------------------------------------ tables
    @staticmethod
    def _borders(table, val='single', color='000000', sz='4'):
        tblPr = table._tbl.tblPr
        b = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement(f'w:{edge}')
            e.set(qn('w:val'), val)
            e.set(qn('w:sz'), sz)
            e.set(qn('w:space'), '0')
            e.set(qn('w:color'), color)
            b.append(e)
        tblPr.insert_element_before(b, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption',
                                    'w:tblDescription', 'w:tblPrChange')

    def table(self, caption, header, rows, widths, size=10.5, center_cols=(), header_fill='DCE6F1', count=False):
        self.ntab += 1
        tot = sum(widths)
        if tot > self.TEXT_W:
            widths = [w * self.TEXT_W / tot for w in widths]
        save = self._count_key
        self._count_key = None
        self.para(f'表{self.part}-{self.ntab}  {caption}', size=10.5, align='center', indent=False, before=4, after=2,
                  line=1.25, keep_next=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._borders(t)
        tblPr = t._tbl.tblPr
        # t.autofit = False already writes <w:tblLayout w:type="fixed"/>
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(w * 567)))
        for ri, row in enumerate([header] + rows):
            tr = t.rows[ri]
            if ri == 0:
                trPr = tr._tr.get_or_add_trPr()
                th = OxmlElement('w:tblHeader')
                trPr.append(th)
            for ci, txt in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Cm(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri == 0 and header_fill:
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), header_fill)
                    tcPr.insert_element_before(shd, 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText',
                                               'w:vAlign', 'w:hideMark', 'w:headers', 'w:cellIns', 'w:cellDel',
                                               'w:cellMerge', 'w:tcPrChange')
                lines = str(txt).split('\n')
                for li, line in enumerate(lines):
                    cp = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    self._nogrid(cp)
                    cp.paragraph_format.line_spacing = 1.15
                    cp.paragraph_format.space_after = Pt(0)
                    cp.alignment = ALIGN['center'] if (ri == 0 or ci in center_cols) else ALIGN['left']
                    self._count_key = save if (count and ri > 0) else None
                    self.runs(cp, line, size=size, bold=(ri == 0))
                    self._count_key = None
        sp = self.doc.add_paragraph()
        sp.paragraph_format.line_spacing = 0.8
        sp.paragraph_format.space_after = Pt(2)
        self._count_key = save
        return f'表{self.part}-{self.ntab}'

    def save(self, path):
        self.doc.save(path)
